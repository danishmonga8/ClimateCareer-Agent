"""CV text extraction for validated PDF and DOCX documents."""

from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document

from app.services.document_service import (
    DocumentValidationError,
    ValidatedDocument,
    validate_document,
)


class DocumentParsingError(ValueError):
    """Raised when text cannot be safely extracted from a document."""


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    """Text and metadata extracted from a validated document."""

    filename: str
    text: str
    page_count: int
    sha256_digest: str


def _clean_extracted_text(text: str) -> str:
    """Remove null characters and unnecessary surrounding whitespace."""
    cleaned_lines = [line.strip() for line in text.replace("\x00", "").splitlines() if line.strip()]
    return "\n".join(cleaned_lines)


def _extract_pdf_text(document: ValidatedDocument) -> tuple[str, int]:
    """Extract text from a non-encrypted PDF."""
    try:
        with fitz.open(document.path) as pdf:
            if pdf.needs_pass:
                raise DocumentParsingError("Encrypted PDFs are not supported.")

            page_text = [page.get_text("text") for page in pdf]
            return "\n".join(page_text), pdf.page_count
    except DocumentParsingError:
        raise
    except Exception as error:
        raise DocumentParsingError("Unable to read the PDF document.") from error


def _extract_docx_text(document: ValidatedDocument) -> tuple[str, int]:
    """Extract paragraphs and table content from a DOCX document."""
    try:
        docx_file = Document(document.path)
        content: list[str] = []

        content.extend(paragraph.text for paragraph in docx_file.paragraphs)

        for table in docx_file.tables:
            for row in table.rows:
                content.append(" | ".join(cell.text for cell in row.cells))

        return "\n".join(content), 1
    except Exception as error:
        raise DocumentParsingError("Unable to read the DOCX document.") from error


def parse_cv(document_path: str | Path) -> ParsedDocument:
    """Validate a CV and extract its readable text."""
    try:
        document = validate_document(document_path)
    except DocumentValidationError as error:
        raise DocumentParsingError(str(error)) from error

    if document.extension == ".pdf":
        raw_text, page_count = _extract_pdf_text(document)
    elif document.extension == ".docx":
        raw_text, page_count = _extract_docx_text(document)
    else:
        raise DocumentParsingError(f"Unsupported document type: {document.extension}")

    cleaned_text = _clean_extracted_text(raw_text)

    if not cleaned_text:
        raise DocumentParsingError("No readable text was found. The document may be image-based.")

    return ParsedDocument(
        filename=document.filename,
        text=cleaned_text,
        page_count=page_count,
        sha256_digest=document.sha256_digest,
    )
